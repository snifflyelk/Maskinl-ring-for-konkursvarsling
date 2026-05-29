from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

DOC_PATH = Path(r"c:\Users\olavb\OneDrive\Documents\Masteroppgave\Masteroppgave Magnus og Olav v 24052026 - sannhetstilpasset-med-figurer.docx")

FIG_ROOT = Path(r"c:\Users\olavb\OneDrive\Documents\Masteroppgave\figurer_kap3_5")

# (heading anchor, image path, image width)
FIGURE_PLAN = [
    ("3.1 Forskningsdesign", FIG_ROOT / "kap3_metode" / "figur_3_1_analysepipeline.png", Inches(6.5)),
    ("3.3.1 Variabler brukt i modellene", FIG_ROOT / "kap3_metode" / "figur_3_2_variabeloversikt.png", Inches(6.5)),
    ("3.2 Datagrunnlag", FIG_ROOT / "kap3_metode" / "figur_3_3_klasseubalanse.png", Inches(5.8)),
    ("4.3.3 Analyse av fordelinger", FIG_ROOT / "kap4_eda" / "figur_4_1_fordeling_alder.png", Inches(6.2)),
    ("4.3.3 Analyse av fordelinger", FIG_ROOT / "kap4_eda" / "figur_4_2_fordeling_ansatte.png", Inches(6.2)),
    ("4.3.4 Analyse av forskjeller mellom konkursselskaper og ikke-konkursselskaper", FIG_ROOT / "kap4_eda" / "figur_4_3_konkursrate_bransje.png", Inches(6.5)),
    ("4.3.4 Analyse av forskjeller mellom konkursselskaper og ikke-konkursselskaper", FIG_ROOT / "kap4_eda" / "figur_4_4_konkursrate_fylke.png", Inches(6.5)),
    ("4.3.5 Korrelasjonsanalyse", FIG_ROOT / "kap4_eda" / "figur_4_5_korrelasjonsmatrise.png", Inches(6.5)),
    ("5.1.3 Modellens prediksjonsevne", FIG_ROOT / "kap5_resultater" / "figur_5_1_roc_sammenligning.png", Inches(6.2)),
    ("5.1.3 Modellens prediksjonsevne", FIG_ROOT / "kap5_resultater" / "figur_5_2_pr_sammenligning.png", Inches(6.2)),
    ("5.3.1 Endringer i modellprestasjon ved ulike terskler", FIG_ROOT / "kap5_resultater" / "figur_5_3_terskelanalyse.png", Inches(6.5)),
    ("5.2.2 Modellens prediksjonsevne", FIG_ROOT / "kap5_resultater" / "figur_5_4_confusion_matrices.png", Inches(6.5)),
    ("5.2.1 Variabelbetydning (Feature Importance)", FIG_ROOT / "kap5_resultater" / "figur_5_5_feature_importance_xgb.png", Inches(6.2)),
]


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style is not None:
        new_para.style = style
    return new_para


def find_first_paragraph(doc: Document, exact_text: str) -> Paragraph | None:
    for p in doc.paragraphs:
        if p.text.strip() == exact_text:
            return p
    return None


def main() -> None:
    doc = Document(str(DOC_PATH))

    inserted = 0
    missing_anchors = []
    missing_files = []

    for anchor, fig_path, fig_width in FIGURE_PLAN:
        if not fig_path.exists():
            missing_files.append(str(fig_path))
            continue

        anchor_p = find_first_paragraph(doc, anchor)
        if anchor_p is None:
            missing_anchors.append(anchor)
            continue

        pic_p = insert_paragraph_after(anchor_p)
        pic_p.alignment = 1  # center
        run = pic_p.add_run()
        run.add_picture(str(fig_path), width=fig_width)
        inserted += 1

    doc.save(str(DOC_PATH))

    print(f"Inserted figures: {inserted}")
    if missing_anchors:
        print("Missing anchors:")
        for a in missing_anchors:
            print(f" - {a}")
    if missing_files:
        print("Missing figure files:")
        for f in missing_files:
            print(f" - {f}")


if __name__ == "__main__":
    main()
